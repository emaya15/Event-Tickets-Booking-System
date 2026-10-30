"""Reports and dashboard. Every number is computed from real database rows.

Money definitions (all in rupees):
* collected = sum of payments whose status is Success or Refunded (Failed payments are never revenue,
  and the unique (booking, success_key) constraint prevents a booking from being counted twice)
* refunded  = sum of refund records
* net       = collected - refunded
"""
import io
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import and_, func, select
from sqlalchemy.orm import Session

from app.models import BookedSeat, Booking, Event, Payment, Refund, Screen, Seat, Show, Ticket, User
from app.utils.constants import BookingStatus, EventStatus, PaymentStatus, Role
from app.utils.errors import BadRequest, Forbidden, NotFound
from app.utils.time import utc_now

COLLECTED = (PaymentStatus.SUCCESS.value, PaymentStatus.REFUNDED.value)
MAX_RANGE_DAYS = 366


def _f(value) -> float:
    return round(float(value or 0), 2)


def resolve_range(from_date: Optional[date], to_date: Optional[date], forward_days: int = 0) -> Tuple[date, date]:
    """Default to the last 30 days (plus ``forward_days`` ahead); validate order and maximum span."""
    today = utc_now().date()
    to_date = to_date or (today + timedelta(days=forward_days))
    from_date = from_date or (min(to_date, today) - timedelta(days=30))
    if from_date > to_date:
        raise BadRequest("from_date must not be after to_date")
    if (to_date - from_date).days > MAX_RANGE_DAYS:
        raise BadRequest(f"The date range may not exceed {MAX_RANGE_DAYS} days")
    return from_date, to_date


def _day_bounds(d1: date, d2: date) -> Tuple[datetime, datetime]:
    return datetime.combine(d1, datetime.min.time()), datetime.combine(d2 + timedelta(days=1), datetime.min.time())


# ------------------------------------------------------------- dashboard
def admin_dashboard(db: Session) -> dict:
    today = utc_now().date()
    start, end = _day_bounds(today, today)
    bookings_today = db.scalar(select(func.count()).select_from(Booking).where(Booking.created_at >= start, Booking.created_at < end))
    confirmed_today = db.scalar(
        select(func.count()).select_from(Booking).where(Booking.confirmed_at >= start, Booking.confirmed_at < end)
    )
    collected = db.scalar(
        select(func.coalesce(func.sum(Payment.amount), 0)).where(
            Payment.status.in_(COLLECTED), Payment.paid_at >= start, Payment.paid_at < end
        )
    )
    refunded = db.scalar(
        select(func.coalesce(func.sum(Refund.amount), 0)).where(Refund.refunded_at >= start, Refund.refunded_at < end)
    )
    refund_count = db.scalar(
        select(func.count()).select_from(Refund).where(Refund.refunded_at >= start, Refund.refunded_at < end)
    )
    active_events = db.scalar(
        select(func.count()).select_from(Event).where(Event.status == EventStatus.PUBLISHED.value)
    )
    scanned = db.scalar(select(func.count()).select_from(Ticket).where(Ticket.used_at >= start, Ticket.used_at < end))
    return {
        "date": today.isoformat(),
        "bookings_today": bookings_today,
        "confirmed_bookings_today": confirmed_today,
        "revenue_collected_today": _f(collected),
        "refunds_today": {"count": refund_count, "amount": _f(refunded)},
        "net_revenue_today": _f(Decimal(collected or 0) - Decimal(refunded or 0)),
        "active_events": active_events,
        "tickets_scanned_today": scanned,
    }


# ----------------------------------------------------------- event sales
def _event_scope(user: User, event_id: Optional[int], db: Session):
    """Returns the event filter for the caller: admins any, organizers only their own events."""
    if user.role == Role.ADMIN.value:
        return event_id
    if user.role == Role.ORGANIZER.value:
        if event_id is not None:
            event = db.get(Event, event_id)
            if event is None:
                raise NotFound("Event not found")
            if event.organizer_id != user.id:
                raise Forbidden("You can only view reports for your own events")
        return event_id
    raise Forbidden("You do not have permission to view this report")


def event_sales(db: Session, user: User, from_date: Optional[date], to_date: Optional[date], event_id: Optional[int]) -> dict:
    event_id = _event_scope(user, event_id, db)
    from_date, to_date = resolve_range(from_date, to_date, forward_days=90)  # shows are often in the future

    show_stmt = (
        select(Show, Event)
        .join(Event, Event.id == Show.event_id)
        .where(Show.show_date >= from_date, Show.show_date <= to_date)
        .order_by(Event.id, Show.start_at)
    )
    if user.role == Role.ORGANIZER.value:
        show_stmt = show_stmt.where(Event.organizer_id == user.id)
    if event_id is not None:
        show_stmt = show_stmt.where(Event.id == event_id)

    events: dict = {}
    for show, event in db.execute(show_stmt).all():
        tickets = db.scalar(
            select(func.coalesce(func.sum(Booking.seat_count), 0)).where(
                Booking.show_id == show.id, Booking.status == BookingStatus.CONFIRMED.value
            )
        )
        collected = db.scalar(
            select(func.coalesce(func.sum(Payment.amount), 0))
            .join(Booking, Booking.id == Payment.booking_id)
            .where(Booking.show_id == show.id, Payment.status.in_(COLLECTED))
        )
        refunded = db.scalar(
            select(func.coalesce(func.sum(Refund.amount), 0))
            .join(Booking, Booking.id == Refund.booking_id)
            .where(Booking.show_id == show.id)
        )
        row = {
            "show_id": show.id,
            "show_start_at": show.start_at.isoformat() + "Z",
            "show_status": show.status,
            "tickets_sold": int(tickets),
            "revenue_collected": _f(collected),
            "refunds": _f(refunded),
            "net_revenue": _f(Decimal(collected or 0) - Decimal(refunded or 0)),
        }
        entry = events.setdefault(
            event.id,
            {"event_id": event.id, "title": event.title, "tickets_sold": 0, "revenue_collected": 0.0, "refunds": 0.0, "net_revenue": 0.0, "shows": []},
        )
        entry["shows"].append(row)
        for key in ("tickets_sold", "revenue_collected", "refunds", "net_revenue"):
            entry[key] = round(entry[key] + row[key], 2) if key != "tickets_sold" else entry[key] + row[key]

    items = list(events.values())
    return {
        "from_date": from_date.isoformat(),
        "to_date": to_date.isoformat(),
        "events": items,
        "totals": {
            "tickets_sold": sum(e["tickets_sold"] for e in items),
            "revenue_collected": round(sum(e["revenue_collected"] for e in items), 2),
            "refunds": round(sum(e["refunds"] for e in items), 2),
            "net_revenue": round(sum(e["net_revenue"] for e in items), 2),
        },
    }


def event_sales_excel(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Event sales"
    ws.append([f"Event sales {report['from_date']} to {report['to_date']}"])
    ws["A1"].font = Font(bold=True, size=13)
    headers = ["Event ID", "Event", "Show ID", "Show start (UTC)", "Show status", "Tickets sold", "Revenue collected", "Refunds", "Net revenue"]
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for event in report["events"]:
        for show in event["shows"]:
            ws.append(
                [event["event_id"], event["title"], show["show_id"], show["show_start_at"], show["show_status"],
                 show["tickets_sold"], show["revenue_collected"], show["refunds"], show["net_revenue"]]
            )
    totals = report["totals"]
    ws.append([])
    ws.append(["", "TOTAL", "", "", "", totals["tickets_sold"], totals["revenue_collected"], totals["refunds"], totals["net_revenue"]])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHI", (10, 32, 9, 22, 13, 13, 18, 12, 14)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ------------------------------------------------------------- occupancy
def occupancy(db: Session, user: User, from_date: Optional[date], to_date: Optional[date], event_id: Optional[int]) -> dict:
    event_id = _event_scope(user, event_id, db)
    from_date, to_date = resolve_range(from_date, to_date, forward_days=90)
    stmt = (
        select(Show, Event, Screen)
        .join(Event, Event.id == Show.event_id)
        .join(Screen, Screen.id == Show.screen_id)
        .where(Show.show_date >= from_date, Show.show_date <= to_date)
        .order_by(Show.start_at)
    )
    if user.role == Role.ORGANIZER.value:
        stmt = stmt.where(Event.organizer_id == user.id)
    if event_id is not None:
        stmt = stmt.where(Event.id == event_id)

    shows = []
    for show, event, screen in db.execute(stmt).all():
        seat_cats = dict(
            db.execute(
                select(Seat.category, func.count()).where(Seat.screen_id == screen.id, Seat.is_active.is_(True)).group_by(Seat.category)
            ).all()
        )
        booked_cats = dict(
            db.execute(
                select(Seat.category, func.count())
                .join(BookedSeat, BookedSeat.seat_id == Seat.id)
                .join(Booking, Booking.id == BookedSeat.booking_id)
                .where(BookedSeat.show_id == show.id, Booking.status == BookingStatus.CONFIRMED.value)
                .group_by(Seat.category)
            ).all()
        )
        total = sum(seat_cats.values())
        booked = sum(booked_cats.values())
        shows.append(
            {
                "show_id": show.id,
                "event_id": event.id,
                "event_title": event.title,
                "show_start_at": show.start_at.isoformat() + "Z",
                "total_seats": total,
                "booked_seats": booked,
                "occupancy_percent": round(booked * 100 / total, 2) if total else 0.0,
                "by_category": {
                    cat: {
                        "total_seats": count,
                        "booked_seats": booked_cats.get(cat, 0),
                        "occupancy_percent": round(booked_cats.get(cat, 0) * 100 / count, 2) if count else 0.0,
                    }
                    for cat, count in sorted(seat_cats.items())
                },
            }
        )
    return {"from_date": from_date.isoformat(), "to_date": to_date.isoformat(), "shows": shows}


# ------------------------------------------------------------ top events
def top_events(db: Session, from_date: Optional[date], to_date: Optional[date]) -> dict:
    from_date, to_date = resolve_range(from_date, to_date)
    start, end = _day_bounds(from_date, to_date)
    rows = db.execute(
        select(Event.id, Event.title, func.sum(Booking.seat_count).label("tickets"), func.count(Booking.id).label("bookings"))
        .join(Show, Show.event_id == Event.id)
        .join(Booking, Booking.show_id == Show.id)
        .where(Booking.status == BookingStatus.CONFIRMED.value, Booking.confirmed_at >= start, Booking.confirmed_at < end)
        .group_by(Event.id, Event.title)
        .order_by(func.sum(Booking.seat_count).desc(), Event.id)
        .limit(10)
    ).all()
    return {
        "from_date": from_date.isoformat(),
        "to_date": to_date.isoformat(),
        "events": [
            {"rank": i, "event_id": r[0], "title": r[1], "tickets_sold": int(r[2]), "bookings": int(r[3])}
            for i, r in enumerate(rows, start=1)
        ],
    }


# --------------------------------------------------------------- refunds
def refunds_report(db: Session, from_date: Optional[date], to_date: Optional[date]) -> dict:
    from_date, to_date = resolve_range(from_date, to_date)
    start, end = _day_bounds(from_date, to_date)
    rows = db.execute(
        select(Refund.reason, func.count(), func.coalesce(func.sum(Refund.amount), 0))
        .where(Refund.refunded_at >= start, Refund.refunded_at < end)
        .group_by(Refund.reason)
        .order_by(Refund.reason)
    ).all()
    groups = [{"reason": r[0], "count": int(r[1]), "amount": _f(r[2])} for r in rows]
    return {
        "from_date": from_date.isoformat(),
        "to_date": to_date.isoformat(),
        "by_reason": groups,
        "total_count": sum(g["count"] for g in groups),
        "total_amount": round(sum(g["amount"] for g in groups), 2),
    }
